"""Fill the Veles Hack submission template for TangleLens (python-pptx 1.0.2).
Usage: python -I fill_template.py <template.pptx> <screenshots_dir> <out.pptx>"""
import copy
import sys

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

TEMPLATE, SHOTS, OUT = sys.argv[1:4]
REPO = "https://github.com/moe-a-m/tanglelens"
prs = Presentation(TEMPLATE)
s_cover, s_repo, s_summary, s_high = prs.slides


def ph(slide, ptype_name):
    for sh in slide.placeholders:
        if sh.placeholder_format.type is not None and sh.placeholder_format.type.name == ptype_name:
            return sh
    raise KeyError(ptype_name)


def set_text(shape, text):
    """Replace the first run's text, keeping its formatting; drop other runs/paragraphs."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.add_run().text = text


def fill_bullets(shape, items, size=None):
    """items: list of (label, text[, link]) -> one bulleted paragraph each, label in bold.
    Bullets and spacing are inherited from the layout (template paragraph properties are cloned)."""
    tf = shape.text_frame
    template_pPr = tf.paragraphs[0]._p.pPr
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for item in items:
        label, text = item[0], item[1]
        link = item[2] if len(item) > 2 else None
        p = tf.add_paragraph()
        if template_pPr is not None:
            p._p.insert(0, copy.deepcopy(template_pPr))
        if label:
            r = p.add_run()
            r.text = label + " "
            r.font.bold = True
            if size:
                r.font.size = Pt(size)
        r = p.add_run()
        r.text = text
        if size:
            r.font.size = Pt(size)
        if link:
            r.hyperlink.address = link


def add_shot(slide, path, left, top, width, alt):
    pic = slide.shapes.add_picture(path, left, top, width=width)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.line.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xD3, 0xDB, 0xE3)
    pic.line.width = Pt(0.75)
    return pic


# --- cover
set_text(ph(s_cover, "TITLE"), "TangleLens")
set_text(ph(s_cover, "SUBTITLE"), "IOTA Advanced Explorer for Eclipse aeriOS · Challenge 2 (O-CEI)")

# --- GitHub repo: text left, UI screenshot right
body = ph(s_repo, "BODY")
body.width = Inches(4.55)
fill_bullets(body, [
    ("GitHub repo:", REPO, REPO),
    ("Licence:", "Apache-2.0, including modified Eclipse aeriOS code (see NOTICE)"),
    ("Run it:", "start the aeriOS tangle (iota-tangle/), then make up · make demo · make test"),
    ("Evidence:", "raw Hornet responses, test summaries and clean-clone runs in reports/"),
    ("Stack:", "Python, FastAPI, Flask, PostgreSQL, Mosquitto MQTT, HORNET 2.0.2, Docker Compose"),
], size=13)
add_shot(s_repo, f"{SHOTS}/1_confirmed_message.png", Inches(5.05), Inches(1.40), Inches(4.6),
         "TangleLens web UI: search filters, message list and a confirmed message with its verification checklist")

# --- Summary: full width
fill_bullets(ph(s_summary, "BODY"), [
    ("What:", "a searchable, continuously verified copy of every message Eclipse aeriOS writes to its private IOTA Tangle. The Tangle stays the authority."),
    ("App 1, extended aeriOS Messages API:", "same /upload contract and HTTP 200; after Hornet accepts a block it forwards the block id and the exact bytes sent, over HTTP and MQTT."),
    ("App 2, Advanced Explorer:", "PostgreSQL copy enriched with readable metadata; REST API and web UI search by block id, date and tag, plus source, text, status and trace."),
    ("Validation:", "every block is checked with Hornet's GET block metadata (solid, milestone-referenced, conflicting) and GET block (content compared as hex, decoded JSON and SHA-256)."),
    ("Audit:", "append-only history of every check; all messages re-audited every 5 minutes, so later tampering is caught."),
    ("Runs on the real aeriOS private tangle", "(HORNET 2.0.2) with Docker Compose; tested from a fresh clone."),
], size=13)

# --- Highlights: text left, tamper screenshot right
body = ph(s_high, "BODY")
body.width = Inches(4.55)
fill_bullets(body, [
    ("Tampering is caught:", "an IE's trust score inflated from 0.61 to 0.95 in the database is flagged, shown next to the Tangle copy and alerted."),
    ("Solid is not confirmed:", "8 explicit states; measured on the real node: solid within 20 ms, milestone median 3.3 s."),
    ("Beyond the brief (UPV ideas #1, #3, #4):", "trace timelines, alerts (UI, webhook, MQTT), MQTT live feed."),
    ("Survives outages:", "explorer or database down for 25 s, no record lost."),
    ("Evidence first:", "115 automated tests (SQLite + PostgreSQL), live smoke test, every Hornet behaviour backed by a saved raw response."),
], size=12)
add_shot(s_high, f"{SHOTS}/2_tampered_message.png", Inches(5.05), Inches(1.40), Inches(4.6),
         "Content mismatch: stored copy with trust_score 0.95 beside the Tangle copy with 0.61; fields that differ: trust_score")

prs.save(OUT)
print("saved", OUT)
